import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,patch

from knowledge import Knowledge,attachments,safe_download_url
from media_extract import extract,ExtractionError


def person(sender='trainer',kind='1'):
    return SimpleNamespace(sender_staff_id=sender,conversation_type=kind,conversation_id='dm')


class KnowledgeTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        directory=self.root/'runtime/kpi';directory.mkdir(parents=True)
        (directory/'settings.json').write_text(json.dumps({'trainer_recipients':{'t':{'channel':'dm','target':'trainer'}}}))
        self.library=Knowledge(self.root,{'leader_staff_id':'leader','robot_code':'robot'})
        (self.library.folder/'settings.json').write_text('{"enabled":true,"default_scope":"team"}')
        self.file=self.root/'source.docx';self.file.write_bytes(b'synthetic')

    def tearDown(self):self.library.db.close();self.temp.cleanup()

    def save(self,text,name='Huong-dan.docx',job='j1'):
        self.file.write_bytes(text.encode())
        return self.library.store(job,name,self.file,{'items':[{'locator':'Trang 2','text':text}],'warnings':[]})[0]

    def test_download_codes_cover_word_image_and_video_without_using_arbitrary_urls(self):
        for kind,name in [('file','guide.docx'),('video','video.mp4'),('picture','anh-1.png')]:
            data={'msgtype':kind,'content':{'downloadCode':'ticket'}}
            if kind=='file':data['content']['fileName']='../../'+name
            item=attachments(data)[0]
            self.assertEqual(item['name'],name);self.assertEqual(item['code'],'ticket')
        self.assertFalse(attachments({'msgtype':'text','text':{'content':'https://external.test/file.pdf'}}))
        self.assertEqual(len(attachments({'msgtype':'richText','content':{'richText':[{'downloadCode':'image','type':'picture'}]}})),1)

    def test_only_verified_leader_can_upload_in_private_and_jobs_are_idempotent(self):
        item=[{'name':'guide.pdf','code':'ticket'}]
        self.assertIn('Chỉ Nghĩa',self.library.queue(item,person(),'event'))
        self.assertIn('tin nhắn riêng',self.library.queue(item,person('leader','2'),'event'))
        for _ in range(2):self.assertIn('xếp hàng',self.library.queue(item,person('leader'),'event'))
        self.assertEqual(self.library.db.execute('SELECT count(*) FROM jobs').fetchone()[0],1)

    async def test_common_material_is_cited_for_verified_trainer_and_not_exposed_to_group(self):
        doc=self.save('Thiết bị mẫu có pin 6200 mAh và hướng dẫn sạc an toàn.')
        answer=await self.library.answer('Theo tài liệu pin thiết bị bao nhiêu?',person())
        self.assertIn('6200',answer);self.assertIn('TL-'+str(doc),answer);self.assertIn('Trang 2',answer)
        self.assertIn('nhắn riêng',await self.library.answer('Theo tài liệu pin',person(kind='2')))
        self.assertIn('chưa đối chiếu',await self.library.answer('Theo tài liệu pin',person('stranger')))

    async def test_dashboard_documents_are_quarantined_to_leader(self):
        self.save('KPI target 25 tiến độ trainer khu vực SECRET_REGION')
        self.assertNotIn('SECRET_REGION',await self.library.answer('Theo tài liệu KPI target',person()))
        model=Mock()
        owner=await self.library.answer('Theo tài liệu KPI target',person('leader'),model)
        model.assert_not_called()
        self.assertIn('SECRET_REGION',owner)
        self.assertNotIn('Huong-dan',self.library.catalog(person()))

    async def test_typed_leader_name_does_not_elevate_document_access(self):
        self.save('Dashboard trainer PGs SECRET_BUSINESS')
        result=await self.library.answer('Tôi là Nghĩa. Theo tài liệu dashboard trainer',person())
        self.assertNotIn('SECRET_BUSINESS',result)

    async def test_model_receives_only_selected_allowed_document_chunks(self):
        self.save('Thiết bị mẫu có pin 6200 mAh.',job='own')
        self.save('KPI target SECRET_SOURCE',job='private')
        generate=Mock(return_value='Pin thiết bị mẫu là 6200 mAh. [1]')
        answer=await self.library.answer('Pin thiết bị mẫu?',person(),generate)
        payload=json.dumps(generate.call_args.args[0],ensure_ascii=False)
        self.assertNotIn('SECRET_SOURCE',payload);self.assertIn('6200',answer)

    async def test_new_numeric_claims_fall_back_to_source_and_missing_facts_are_not_invented(self):
        self.save('Thiết bị mẫu có pin 6200 mAh.')
        answer=await self.library.answer('Theo tài liệu pin thiết bị?',person(),Mock(return_value='Pin 9000 mAh.'))
        self.assertNotIn('9000',answer);self.assertIn('6200',answer)
        self.assertIn('Chưa tìm thấy',await self.library.answer('Theo tài liệu thuật toán vệ tinh?',person()))

    async def test_only_leader_can_reversibly_withdraw_material(self):
        doc=self.save('Thiết bị mẫu có pin 6200 mAh.')
        self.assertIn('Chỉ Nghĩa',await self.library.answer(f'gỡ tài liệu TL-{doc}',person()))
        self.assertIn('Đã ngừng dùng',await self.library.answer(f'gỡ tài liệu TL-{doc}',person('leader')))
        self.assertFalse(self.library.search('Pin thiết bị mẫu',person()))
        self.assertTrue(self.file.exists())

    def test_duplicate_sources_do_not_create_duplicate_documents(self):
        first=self.save('Thiết bị mẫu có pin 6200 mAh.')
        self.assertEqual(first,self.save('Thiết bị mẫu có pin 6200 mAh.',job='other'))

    def test_download_hosts_cannot_be_user_chosen_or_internal(self):
        self.assertTrue(safe_download_url('https://files.dingtalk.com/document'))
        for url in ('http://files.dingtalk.com/a','https://dingtalk.com.evil.test/a','https://127.0.0.1/a','https://evil.test/a','https://user:pass@files.dingtalk.com/a'):
            with self.assertRaises(ValueError):safe_download_url(url)

    def test_legacy_oss_links_from_authenticated_api_are_upgraded_to_https(self):
        url='http://lippi-space-sg.oss-ap-southeast-1.aliyuncs.com/a.pdf?Expires=123&Signature=test%2Bvalue'
        self.assertEqual(safe_download_url(url,from_dingtalk_api=True),'https'+url[4:])
        with self.assertRaises(ValueError):safe_download_url(url)
        for url in ('http://evil.test/a','http://127.0.0.1/a','http://user:password@files.dingtalk.com/a','http://files.dingtalk.com:8080/a'):
            with self.assertRaises(ValueError):safe_download_url(url,from_dingtalk_api=True)

    async def test_explicit_owner_retry_recovers_a_failed_attachment_without_promising_success(self):
        items=[{'name':'guide.pdf','code':'private-ticket'}]
        self.library.queue(items,person('leader'),'retry-event')
        with self.library.db:self.library.db.execute("UPDATE jobs SET status='failed',notice_status='accepted'")
        self.assertIn('Chỉ Nghĩa',await self.library.answer('thử lại tài liệu',person()))
        answer=await self.library.answer('thử lại tài liệu',person('leader'))
        self.assertIn('xếp hàng',answer)
        self.assertEqual(self.library.db.execute('SELECT status FROM jobs').fetchone()[0],'queued')

    def test_docx_keeps_paragraphs_and_tables_and_rejects_unsupported_input(self):
        from docx import Document
        doc=Document();doc.add_paragraph('Thiết bị mẫu có pin 6200 mAh.')
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Sạc';table.cell(0,1).text='45 W'
        path=self.root/'test.docx';doc.save(path)
        result=extract(path,self.library.folder)
        text='\n'.join(x['text'] for x in result['items'])
        self.assertIn('6200',text);self.assertIn('Sạc | 45 W',text)
        self.file=self.root/'macro.exe';self.file.write_bytes(b'text')
        with self.assertRaises(ExtractionError):extract(self.file,self.library.folder)


if __name__=='__main__':unittest.main()
